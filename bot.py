import os
import sqlite3
import threading
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from telegram import (
    Update, ReplyKeyboardMarkup, InlineKeyboardMarkup, InlineKeyboardButton
)
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    ContextTypes, ConversationHandler, filters
)

TOKEN = os.environ["BOT_TOKEN"]
DB_PATH = os.environ.get("DB_PATH", "navbat.db")
XLSX_PATH = os.environ.get("XLSX_PATH", "navbat.xlsx")
PORT = int(os.environ.get("PORT", "10000"))

MENU = [
    ["➕ Navbat berildi", "💰 Shoferdan olindi"],
    ["📊 Hisobot", "📥 Excel"],
]

CARS = [
    "60059KBA", "60227SBA", "60510KBA", "60501KBA", "60507KBA",
    "60511KBA", "60260SBA", "60515KBA", "60277KBA", "60110QBA",
    "60720SBA", "60545SBA", "60740SBA", "60730SBA", "60474SBA",
    "60373SBA", "60414SBA", "60755SBA", "60441SBA", "60442SBA",
    "60445SBA"
]

class HealthHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.end_headers()
        self.wfile.write(b"OK")
    def log_message(self, format, *args):
        return

def start_health_server():
    server = ThreadingHTTPServer(("0.0.0.0", PORT), HealthHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

def db():
    con = sqlite3.connect(DB_PATH)
    con.execute("""
        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            given_date TEXT NOT NULL,
            car TEXT NOT NULL,
            given_amount REAL NOT NULL,
            received_date TEXT,
            received_amount REAL
        )
    """)
    con.commit()
    return con

def fmt(n):
    return f"{n:,.2f}".rstrip("0").rstrip(".")

def export_excel():
    con = db()
    rows = con.execute("""
        SELECT given_date, car, given_amount, received_date, received_amount
        FROM payments
        ORDER BY id
    """).fetchall()
    con.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Navbat"

    headers = ["Berilgan sana", "Mashina", "Berilgan pul",
               "Olingan sana", "Olingan pul", "Farqi"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for given_date, car, given_amount, received_date, received_amount in rows:
        diff = None if received_amount is None else received_amount - given_amount
        ws.append([given_date, car, given_amount, received_date, received_amount, diff])

    for row in ws.iter_rows(min_row=2):
        row[2].number_format = '#,##0.00'
        row[4].number_format = '#,##0.00'
        row[5].number_format = '#,##0.00'

    widths = [16, 16, 16, 16, 16, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(XLSX_PATH)

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Elektron navbat hisob-kitob botiga xush kelibsiz.",
        reply_markup=ReplyKeyboardMarkup(MENU, resize_keyboard=True)
    )

async def date_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    buttons = [
        [
            InlineKeyboardButton("📅 Bugun", callback_data="date:today"),
            InlineKeyboardButton("📅 Kecha", callback_data="date:yesterday"),
        ],
        [InlineKeyboardButton("🗓 Boshqa sana", callback_data="date:other")]
    ]
    await update.message.reply_text(
        "📅 Sanani tanlang:",
        reply_markup=InlineKeyboardMarkup(buttons)
    )
    return 1

async def choose_given(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    context.user_data["mode"] = "given"
    return await date_menu(update, context)

async def choose_received(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    context.user_data["mode"] = "received"
    return await date_menu(update, context)

def car_buttons(mode):
    return [
        [InlineKeyboardButton(CARS[i], callback_data=f"{mode}_car:{CARS[i]}")]
        for i in range(0, len(CARS))
    ]

async def show_cars(query, context):
    mode = context.user_data["mode"]
    await query.edit_message_text(
        f"📅 Sana: {context.user_data['date']}\n🚛 Mashinani tanlang:",
        reply_markup=InlineKeyboardMarkup(car_buttons(mode))
    )

async def date_choice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    choice = q.data.split(":", 1)[1]

    if choice == "today":
        context.user_data["date"] = datetime.now().strftime("%d.%m.%Y")
        await show_cars(q, context)
        return 2
    if choice == "yesterday":
        context.user_data["date"] = (datetime.now() - timedelta(days=1)).strftime("%d.%m.%Y")
        await show_cars(q, context)
        return 2

    await q.edit_message_text("🗓 Sanani kiriting (KK.OO.YYYY):\nMasalan: 29.08.2026")
    return 4

async def date_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    try:
        dt = datetime.strptime(text, "%d.%m.%Y")
    except ValueError:
        await update.message.reply_text("Sana noto'g'ri. Masalan: 29.08.2026")
        return 4

    context.user_data["date"] = dt.strftime("%d.%m.%Y")
    mode = context.user_data["mode"]
    await update.message.reply_text(
        f"📅 Sana: {context.user_data['date']}\n🚛 Mashinani tanlang:",
        reply_markup=InlineKeyboardMarkup(car_buttons(mode))
    )
    return 2

async def car_choice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    mode, car = q.data.split(":", 1)
    context.user_data["car"] = car
    context.user_data["mode"] = "given" if mode.startswith("given") else "received"
    await q.edit_message_text(
        f"🚛 Mashina: {car}\n💰 Summani kiriting:"
    )
    return 3

async def amount_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        amount = float(update.message.text.replace(" ", "").replace(",", "."))
        if amount < 0:
            raise ValueError
    except ValueError:
        await update.message.reply_text("Summani raqamda kiriting.")
        return 3

    mode = context.user_data["mode"]
    date = context.user_data["date"]
    car = context.user_data["car"]

    con = db()

    if mode == "given":
        con.execute(
            "INSERT INTO payments(given_date, car, given_amount) VALUES(?,?,?)",
            (date, car, amount)
        )
        con.commit()
        msg = (
            f"✅ Navbat berildi\n"
            f"📅 {date}\n🚛 {car}\n💰 {fmt(amount)}"
        )
    else:
        row = con.execute(
            """SELECT id, given_amount, given_date FROM payments
               WHERE car=? AND received_amount IS NULL
               ORDER BY id DESC LIMIT 1""",
            (car,)
        ).fetchone()

        if not row:
            con.close()
            await update.message.reply_text(
                f"⚠️ {car} uchun hali olinmagan 'berilgan' yozuv topilmadi."
            )
            return ConversationHandler.END

        con.execute(
            "UPDATE payments SET received_date=?, received_amount=? WHERE id=?",
            (date, amount, row[0])
        )
        con.commit()
        diff = amount - row[1]
        sign = "+" if diff > 0 else ""
        msg = (
            f"✅ Shoferdan olindi\n"
            f"📅 {date}\n🚛 {car}\n💰 {fmt(amount)}\n"
            f"📌 Farqi: {sign}{fmt(diff)}"
        )

    con.close()
    export_excel()
    await update.message.reply_text(msg)
    return ConversationHandler.END

async def report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    con = db()
    rows = con.execute("""
        SELECT car,
               COALESCE(SUM(given_amount),0),
               COALESCE(SUM(received_amount),0),
               COALESCE(SUM(
                   CASE WHEN received_amount IS NOT NULL
                        THEN received_amount - given_amount ELSE 0 END
               ),0),
               SUM(CASE WHEN received_amount IS NULL THEN 1 ELSE 0 END)
        FROM payments
        GROUP BY car
        ORDER BY car
    """).fetchall()

    total_given = con.execute("SELECT COALESCE(SUM(given_amount),0) FROM payments").fetchone()[0]
    total_received = con.execute("SELECT COALESCE(SUM(received_amount),0) FROM payments").fetchone()[0]
    total_diff = con.execute("""
        SELECT COALESCE(SUM(CASE WHEN received_amount IS NOT NULL
        THEN received_amount - given_amount ELSE 0 END),0) FROM payments
    """).fetchone()[0]
    pending = con.execute(
        "SELECT COUNT(*) FROM payments WHERE received_amount IS NULL"
    ).fetchone()[0]
    con.close()

    if not rows:
        await update.message.reply_text("📊 Hozircha ma'lumot yo'q.")
        return

    lines = ["📊 MASHINALAR BO'YICHA HISOBOT", ""]
    for car, given, received, diff, pending_car in rows:
        lines.append(
            f"🚛 {car}\n"
            f"   Berilgan: {fmt(given)}\n"
            f"   Olingan:  {fmt(received)}\n"
            f"   Farqi:    {fmt(diff)}\n"
            f"   Holati:   {'⏳ Olinmagan' if pending_car else '✅ Yopilgan'}\n"
        )

    lines.extend([
        "━━━━━━━━━━━━━━",
        "📌 UMUMIY",
        f"Berilgan jami: {fmt(total_given)}",
        f"Olingan jami:  {fmt(total_received)}",
        f"Farqi:         {fmt(total_diff)}",
        f"Olinmagan yozuvlar: {pending}",
    ])
    await update.message.reply_text("\n".join(lines))

async def send_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    export_excel()
    with open(XLSX_PATH, "rb") as f:
        await update.message.reply_document(
            document=f,
            filename="Navbat_RICHES.xlsx",
            caption="📥 Elektron navbat jadvali"
        )

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("Bekor qilindi.")
    return ConversationHandler.END

def main():
    db().close()
    export_excel()
    start_health_server()

    app = Application.builder().token(TOKEN).build()

    given_conv = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex("^➕ Navbat berildi$"), choose_given)],
        states={
            1: [
                CallbackQueryHandler(date_choice, pattern=r"^date:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, date_input),
            ],
            2: [CallbackQueryHandler(car_choice, pattern=r"^given_car:")],
            3: [MessageHandler(filters.TEXT & ~filters.COMMAND, amount_input)],
            4: [MessageHandler(filters.TEXT & ~filters.COMMAND, date_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel)]
    )

    received_conv = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex("^💰 Shoferdan olindi$"), choose_received)],
        states={
            1: [
                CallbackQueryHandler(date_choice, pattern=r"^date:"),
                MessageHandler(filters.TEXT & ~filters.COMMAND, date_input),
            ],
            2: [CallbackQueryHandler(car_choice, pattern=r"^received_car:")],
            3: [MessageHandler(filters.TEXT & ~filters.COMMAND, amount_input)],
            4: [MessageHandler(filters.TEXT & ~filters.COMMAND, date_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel)]
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(given_conv)
    app.add_handler(received_conv)
    app.add_handler(MessageHandler(filters.Regex("^📊 Hisobot$"), report))
    app.add_handler(MessageHandler(filters.Regex("^📥 Excel$"), send_excel))

    print("Bot started")
    app.run_polling()

if __name__ == "__main__":
    main()
